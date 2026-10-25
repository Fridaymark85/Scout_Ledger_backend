"""SAP-compatible Excel export for approved requisitions (in-memory, openpyxl)."""
from __future__ import annotations

import io
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from .models import Requisition, RequisitionItem

# SAP field names for exactly the values the request form collects (nothing else is exported).
COLUMNS: tuple[str, ...] = (
    "EKGRP", "TXZ01", "MENGE", "MEINS", "NETPR", "WERKS",
    "LGORT", "KOSTL", "EEIND", "MATKL", "AFNAM", "WAERS",
)

# Plain-English meaning of each SAP code (shown as header hover notes and on the Field_Guide sheet).
FIELD_LABELS: dict[str, str] = {
    "EKGRP": "Purchasing group (buyer group)",
    "TXZ01": "Short text / description of the item (max 40 characters)",
    "MENGE": "Quantity requested",
    "MEINS": "Unit of measure (EA, KG, L, ...)",
    "NETPR": "Net price per unit (materials: budgetary total / quantity)",
    "WERKS": "Plant",
    "LGORT": "Storage location",
    "KOSTL": "Cost center (account assignment)",
    "EEIND": "Delivery date (DD.MM.YYYY)",
    "MATKL": "Material group",
    "AFNAM": "Requisitioner / indentor ID",
    "WAERS": "Currency",
}

# Leading characters that spreadsheets may interpret as a formula
# (the four required ones plus tab/CR, which OWASP also lists).
_DANGEROUS_PREFIXES: tuple[str, ...] = ("=", "+", "-", "@", "\t", "\r")


def sanitize_cell(value: Any) -> Any:
    """Neutralise formula injection: prefix risky strings with a single quote.

    Non-string values (Decimal/int) are returned untouched so numbers stay numeric.
    """
    if isinstance(value, str) and value.startswith(_DANGEROUS_PREFIXES):
        return "'" + value
    return value


def _unit_price(item: RequisitionItem) -> Decimal:
    """Services carry a unit gross price; materials only a budgetary total, so derive the unit price."""
    if item.item_type == RequisitionItem.ItemType.SERVICE:
        return item.gross_price or Decimal("0")
    return ((item.total_value or Decimal("0")) / item.quantity).quantize(Decimal("0.01"))


def generate_requisition_workbook(requisition: Requisition) -> io.BytesIO:
    """Build an .xlsx with one row per line item of an APPROVED requisition."""
    if requisition.status != Requisition.Status.APPROVED:
        raise ValueError("Only APPROVED requisitions can be exported.")

    wb = Workbook()
    ws = wb.active
    ws.title = "PR_UPLOAD"

    ws.append(list(COLUMNS))
    for cell in ws[1]:
        cell.font = Font(bold=True)
        # Hover note in Excel with the plain-English meaning of the SAP code.
        cell.comment = Comment(FIELD_LABELS[cell.value], "Scout Ledger")

    for item in requisition.items.all():
        row = [
            item.purchasing_group, item.short_text,
            item.quantity, item.unit_of_measure, _unit_price(item), item.plant,
            item.storage_location, item.cost_center,
            item.delivery_date.strftime("%d.%m.%Y") if item.delivery_date else "",
            item.material_group, item.requisitioner, item.currency,
        ]
        ws.append([sanitize_cell(v) for v in row])

    for col in ws.columns:  # readable column widths
        ws.column_dimensions[col[0].column_letter].width = 16

    # Second sheet explaining every SAP code; the upload sheet above stays untouched.
    guide = wb.create_sheet("Field_Guide")
    guide.append(["SAP field", "Meaning"])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    for code in COLUMNS:
        guide.append([code, FIELD_LABELS[code]])
    guide.column_dimensions["A"].width = 12
    guide.column_dimensions["B"].width = 70

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
